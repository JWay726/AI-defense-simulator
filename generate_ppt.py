"""
根据项目代码生成精简版答辩 PPT（6 页，5-7 分钟讲解）
结构：需求分析 / 系统架构 / 商业分析
运行：python generate_ppt.py
"""
import os
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# ============================================================
# 配置
# ============================================================
PROJECT_DIR = r"D:\NCUT\小学期\AI答辩"
OUTPUT_PPT = os.path.join(PROJECT_DIR, "答辩PPT_AI答辩模拟器.pptx")
SUBSCRIPTION_IMG = r"C:\Users\23895\.workbuddy\clipboard-images\clipboard-2026-07-24T02-29-35-026Z-5693f1c6.png"
DIAGRAM_ARCH = os.path.join(PROJECT_DIR, "diagram_architecture.png")
DIAGRAM_FLOW = os.path.join(PROJECT_DIR, "diagram_flow.png")

C_INDOGO = RGBColor(79, 70, 229)
C_INDOGO_LIGHT = RGBColor(238, 242, 255)
C_GRAY = RGBColor(75, 85, 99)
C_DARK = RGBColor(17, 24, 39)
C_WHITE = RGBColor(255, 255, 255)
C_ROSE = RGBColor(225, 29, 72)
C_AMBER = RGBColor(245, 158, 11)
C_EMERALD = RGBColor(16, 185, 129)

FONT_NAME = "Microsoft YaHei"


def get_font(size):
    candidates = [
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\msyhbd.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
    ]
    for path in candidates:
        if os.path.exists(path):
            try:
                return ImageFont.truetype(path, size)
            except Exception:
                pass
    return ImageFont.load_default()


def add_textbox(slide, text, left, top, width, height, font_size=20, bold=False,
                color=C_DARK, align=PP_ALIGN.LEFT, font_name=FONT_NAME, line_spacing=1.2):
    tf = slide.shapes.add_textbox(left, top, width, height).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.name = font_name
    p.font.color.rgb = color
    p.alignment = align
    tf.line_spacing = line_spacing
    return tf


def add_title(slide, text, subtitle=None):
    add_textbox(slide, text, Inches(0.5), Inches(0.35), Inches(12.3), Inches(0.8),
                font_size=30, bold=True, color=C_INDOGO)
    if subtitle:
        add_textbox(slide, subtitle, Inches(0.5), Inches(1.05), Inches(12.3), Inches(0.45),
                    font_size=15, color=C_GRAY)
    line = slide.shapes.add_shape(1, Inches(0.5), Inches(1.42), Inches(12.3), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = C_INDOGO
    line.line.color.rgb = C_INDOGO


def add_bullets(slide, bullets, left, top, width, height, font_size=18, color=C_DARK):
    tf = slide.shapes.add_textbox(left, top, width, height).text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"●  {b}"
        p.font.size = Pt(font_size)
        p.font.name = FONT_NAME
        p.font.color.rgb = color
        p.space_before = Pt(8)
        p.line_spacing = 1.25
    return tf


def draw_architecture_diagram(path):
    W, H = 1100, 460
    img = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(img)
    font_title = get_font(20)
    font_text = get_font(16)
    font_small = get_font(14)

    def box(x, y, w, h, title, subs, color):
        r = 12
        draw.rounded_rectangle([x, y, x + w, y + h], radius=r, fill=color, outline="#C7D2FE", width=2)
        draw.text((x + 15, y + 10), title, font=font_title, fill="#1E1B4B")
        ty = y + 44
        for s in subs:
            draw.text((x + 15, ty), s, font=font_small, fill="#4B5563")
            ty += 20

    box(40, 40, 270, 160, "React 前端", [
        "Vite + React 19 + Tailwind",
        "GenUI 卡片式组件流",
        "Web Speech API 语音",
    ], "#EEF2FF")
    box(400, 110, 270, 180, "Node.js 后端", [
        "Express 5 + MySQL2",
        "JWT / bcrypt 认证",
        "文档解析：pdfjs/mammoth",
        "会话与订阅管理",
    ], "#FFF7ED")
    box(760, 110, 270, 180, "DeepSeek API", [
        "文档画像抽取",
        "问题/评分/追问生成",
        "报告与话术模板",
    ], "#F0FDF4")
    box(400, 320, 270, 110, "本地文档解析", [
        "PDF.js 提取 PDF 文本",
        "Mammoth 读 Word",
    ], "#FDF2F8")

    def arrow(x1, y1, x2, y2):
        draw.line([(x1, y1), (x2, y2)], fill="#6366F1", width=3)
        if x1 == x2:
            dy = 8 if y2 > y1 else -8
            draw.polygon([(x1, y2), (x1 - 5, y2 - dy), (x1 + 5, y2 - dy)], fill="#6366F1")
        else:
            dx = 8 if x2 > x1 else -8
            draw.polygon([(x2, y1), (x2 - dx, y1 - 5), (x2 - dx, y1 + 5)], fill="#6366F1")

    arrow(310, 120, 400, 170)
    arrow(670, 170, 760, 170)
    arrow(535, 290, 535, 320)
    arrow(760, 230, 670, 230)
    img.save(path)


def draw_flow_diagram(path):
    W, H = 1200, 230
    img = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(img)
    font = get_font(16)
    font_title = get_font(18)

    steps = [
        ("1 上传文档", "PDF/Word/PPT"),
        ("2 AI 解析", "画像+证据锚点"),
        ("3 选择题型", "动态分类"),
        ("4 语音作答", "Web Speech"),
        ("5 AI 评分", "5 维度"),
        ("6 追问/话术", "致命追问"),
        ("7 综合报告", "复盘建议"),
    ]
    box_w, box_h = 150, 90
    gap = 18
    start_x = 35
    y = 70
    colors = ["#EEF2FF", "#FDF2F8", "#FFF7ED", "#ECFDF5", "#E0F2FE", "#FEF3C7", "#F3E8FF"]
    for i, (title, sub) in enumerate(steps):
        x = start_x + i * (box_w + gap)
        draw.rounded_rectangle([x, y, x + box_w, y + box_h], radius=10, fill=colors[i], outline="#C7D2FE", width=2)
        draw.text((x + 10, y + 12), title, font=font_title, fill="#1E1B4B")
        draw.text((x + 10, y + 42), sub, font=font, fill="#4B5563")
        if i < len(steps) - 1:
            ax = x + box_w + 2
            ay = y + box_h // 2
            draw.line([(ax, ay), (ax + gap - 2, ay)], fill="#6366F1", width=3)
            draw.polygon([(ax + gap - 2, ay), (ax + gap - 10, ay - 5), (ax + gap - 10, ay + 5)], fill="#6366F1")
    img.save(path)


def make_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 封面
    s = prs.slides.add_slide(blank)
    shape = s.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(3.4))
    shape.fill.solid()
    shape.fill.fore_color.rgb = C_INDOGO
    shape.line.fill.background()
    add_textbox(s, "AI 答辩模拟器", Inches(0), Inches(1.1), Inches(13.333), Inches(1.0),
                font_size=52, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    add_textbox(s, "需求分析 · 系统架构 · 商业分析", Inches(0), Inches(2.2), Inches(13.333), Inches(0.6),
                font_size=22, color=C_WHITE, align=PP_ALIGN.CENTER)
    add_textbox(s, "答辩人：________\n指导老师：________\n日期：2026.07", Inches(4), Inches(4.3), Inches(5.333), Inches(2.0),
                font_size=18, color=C_GRAY, align=PP_ALIGN.CENTER)

    # 2 需求分析
    s = prs.slides.add_slide(blank)
    add_title(s, "01 需求分析", "痛点、用户与需求规格")
    left = [
        "答辩问题不可预测，缺真实压力陪练",
        "学生难以发现论文逻辑漏洞与数据缺口",
        "口头作答紧张，表达不结构化",
    ]
    right = [
        "目标用户：本/硕/博论文答辩、项目结题学生",
        "核心功能需求：文档解析、动态出题、语音/文字作答、智能评分、报告复盘",
        "非功能需求：7×24 在线、可定制人格、即时反馈、前后端双模式稳定运行",
    ]
    add_textbox(s, "现实痛点", Inches(0.6), Inches(1.8), Inches(5.8), Inches(0.5),
                font_size=22, bold=True, color=C_ROSE)
    add_bullets(s, left, Inches(0.6), Inches(2.3), Inches(5.8), Inches(3.0), font_size=18)
    add_textbox(s, "用户需求与规格", Inches(6.8), Inches(1.8), Inches(5.8), Inches(0.5),
                font_size=22, bold=True, color=C_INDOGO)
    add_bullets(s, right, Inches(6.8), Inches(2.3), Inches(5.8), Inches(3.8), font_size=17)

    # 3 系统架构
    s = prs.slides.add_slide(blank)
    add_title(s, "02 系统架构", "前后端分离 + 大模型驱动")
    tech = [
        "前端：React 19 + TypeScript + Vite 8 + Tailwind CSS v4",
        "后端：Node.js + Express 5 + MySQL（mysql2）",
        "AI：DeepSeek（deepseek-chat） via OpenAI SDK",
        "文档解析：PDF.js / Mammoth；认证：JWT + bcryptjs",
    ]
    add_bullets(s, tech, Inches(0.5), Inches(1.6), Inches(6.0), Inches(2.0), font_size=15)
    s.shapes.add_picture(DIAGRAM_ARCH, Inches(6.6), Inches(1.6), width=Inches(6.2))

    # 4 核心功能与主流程
    s = prs.slides.add_slide(blank)
    add_title(s, "03 核心功能与主流程", "从文档到报告的一站式演练")
    feats = [
        "文档解析：PDF/Word/PPT/TXT → 结构化画像 + 证据锚点",
        "动态出题：AI 基于文档证据生成 3-5 类追问方向",
        "语音/文字作答：Web Speech API 强制开口，3 分钟倒计时",
        "5 维度评分 + 致命追问 + 标准话术模板",
        "类型复盘与综合报告，输出可执行提升建议",
    ]
    add_bullets(s, feats, Inches(0.5), Inches(1.6), Inches(12.3), Inches(2.1), font_size=16)
    s.shapes.add_picture(DIAGRAM_FLOW, Inches(0.5), Inches(3.8), width=Inches(12.3))

    # 5 商业分析
    s = prs.slides.add_slide(blank)
    add_title(s, "04 商业分析", "订阅策略与价值主张")
    biz = [
        "商业模式：Freemium + 订阅增值",
        "免费层：降低使用门槛，吸引种子用户",
        "付费层：解锁追问、提升题量上限，满足深度演练需求",
        "价值主张：让每个学生都能低成本拥有一位 7×24 的专属答辩教练",
    ]
    add_bullets(s, biz, Inches(0.5), Inches(1.7), Inches(6.2), Inches(4.5), font_size=16)
    if os.path.exists(SUBSCRIPTION_IMG):
        s.shapes.add_picture(SUBSCRIPTION_IMG, Inches(6.9), Inches(1.6), width=Inches(6.0))

    # 6 总结与展望
    s = prs.slides.add_slide(blank)
    add_title(s, "05 总结与展望", "已完成 & 下一步")
    left2 = [
        "MVP 已跑通：登录/解析/出题/评分/追问/报告全闭环",
        "需求分析、系统架构、商业分析三部分完整落地",
        "前后端双模式设计，演示稳定可离线",
    ]
    right2 = [
        "接入 OCR，支持扫描件和图片版 PPT",
        "更多导师人格与多人实时答辩",
        "RAG 跨文档问答、移动端小程序",
    ]
    add_textbox(s, "项目总结", Inches(0.6), Inches(1.8), Inches(5.8), Inches(0.5),
                font_size=20, bold=True, color=C_INDOGO)
    add_bullets(s, left2, Inches(0.6), Inches(2.3), Inches(5.8), Inches(3.5), font_size=17)
    add_textbox(s, "未来展望", Inches(6.8), Inches(1.8), Inches(5.8), Inches(0.5),
                font_size=20, bold=True, color=C_AMBER)
    add_bullets(s, right2, Inches(6.8), Inches(2.3), Inches(5.8), Inches(3.5), font_size=17)
    add_textbox(s, "感谢聆听 · 欢迎提问", Inches(0), Inches(6.2), Inches(13.333), Inches(0.6),
                font_size=22, bold=True, color=C_INDOGO, align=PP_ALIGN.CENTER)

    prs.save(OUTPUT_PPT)
    print(f"已生成：{OUTPUT_PPT}")


if __name__ == "__main__":
    draw_architecture_diagram(DIAGRAM_ARCH)
    draw_flow_diagram(DIAGRAM_FLOW)
    make_presentation()
